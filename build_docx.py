from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / "dist" / "PISA_2025_臺灣學習羅盤.docx"
ASSET = ROOT / "dist" / "oecd-country-note-cover.jpg"
CHARTS = ROOT / "docx_assets"
CHARTS.mkdir(exist_ok=True)

INK = "123044"; BLUE = "087E9A"; CYAN = "54D6D0"; YELLOW = "FFD166"; ORANGE = "FF8A5B"; VIOLET = "8D7DF2"; PALE = "EDF7FA"; LINE = "D9D9D9"

def font(size, bold=False):
    try: return ImageFont.truetype("C:/Windows/Fonts/msjh.ttc", size=size, index=0)
    except Exception: return ImageFont.load_default()

def make_bar_chart(path, title, rows, max_value, suffix=""):
    w, h = 1500, 190 + 115 * len(rows)
    im = Image.new("RGB", (w, h), "white"); d = ImageDraw.Draw(im)
    d.text((70, 35), title, font=font(44, True), fill="#123044")
    left, right = 350, 1370
    for i, (label, value, color, compare) in enumerate(rows):
        y = 140 + i * 115
        d.text((70, y), label, font=font(28, True), fill="#123044")
        d.rounded_rectangle((left, y+4, right, y+42), radius=18, fill="#e8eef1")
        end = left + int((right-left) * value/max_value)
        d.rounded_rectangle((left, y+4, end, y+42), radius=18, fill=color)
        d.text((right+25, y), f"{value}{suffix}", font=font(28, True), fill="#123044")
        if compare is not None:
            y2=y+52; d.rounded_rectangle((left,y2,right,y2+24),radius=12,fill="#eef2f4")
            end2=left+int((right-left)*compare/max_value); d.rounded_rectangle((left,y2,end2,y2+24),radius=12,fill="#aebdc4")
            d.text((right+25,y2-6),f"{compare}{suffix}",font=font(22),fill="#587080")
    im.save(path, quality=92)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        tag = qn(f"w:{edge}"); el = borders.find(tag)
        if el is None: el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), color)

def set_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tcPr.append(mar)
    for k,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = mar.find(qn(f"w:{k}"))
        if node is None: node=OxmlElement(f"w:{k}"); mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def style_run(run, size=11, bold=False, color=INK):
    run.font.name="Microsoft JhengHei"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def add_para(doc, text="", size=11, bold=False, color=INK, align=None, before=0, after=7, keep=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.28
    if align is not None: p.alignment=align
    if keep: p.paragraph_format.keep_with_next=True
    style_run(p.add_run(text),size,bold,color); return p

def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.space_before=Pt(13 if level==1 else 9); p.paragraph_format.space_after=Pt(6); p.paragraph_format.keep_with_next=True
    for r in p.runs: style_run(r,20 if level==1 else 15,True,"000000")
    return p

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; set_cell_shading(c,BLUE); set_cell_border(c); set_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; style_run(p.add_run(h),10,True,"FFFFFF")
        if widths: c.width=Inches(widths[i])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        cant_split = OxmlElement("w:cantSplit")
        t.rows[-1]._tr.get_or_add_trPr().append(cant_split)
        for i,val in enumerate(row):
            c=cells[i]; c.text=""; set_cell_border(c); set_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: set_cell_shading(c,"F1F8FB")
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT; style_run(p.add_run(str(val)),10,False,INK)
            if widths: c.width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(2); return t

make_bar_chart(CHARTS/"scores.png","平均分數：臺灣與 OECD",[("科學 Science",540,"#54d6d0",482),("數學 Mathematics",546,"#ffd166",463),("閱讀 Reading",508,"#ff8a5b",461),("計算問題解決 CPS",551,"#8d7df2",500)],600," 分")
make_bar_chart(CHARTS/"baseline.png","達到基礎熟練度以上的學生比例",[("科學 Science",87,"#54d6d0",74),("數學 Mathematics",85,"#ffd166",65),("閱讀 Reading",82,"#ff8a5b",69)],100,"%")
make_bar_chart(CHARTS/"learning.png","學習投入與自我調節",[("好奇心",73,"#54d6d0",73),("遇難題加倍努力",69,"#31b77a",60),("成長型思維",46,"#8d7df2",69),("規劃讀書方式",30,"#ff8a5b",37),("自我提問檢查理解",33,"#ffd166",49)],100,"%")
make_bar_chart(CHARTS/"environment.png","環境素養與行動",[("環境科學達基準",87,"#31b77a",74),("相信自己能帶來改變",92,"#54d6d0",81),("懂得與他人合作",86,"#ffd166",68),("環境職涯興趣",50,"#ff8a5b",62)],100,"%")

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
for name,size in (("Normal",11),("Title",30),("Heading 1",20),("Heading 2",15)):
    s=styles[name]; s.font.name="Microsoft JhengHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(10); style_run(p.add_run("PISA 2025  |  TAIWAN LEARNING COMPASS"),11,True,BLUE)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); style_run(p.add_run("臺灣學習羅盤"),34,True,"000000")
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(20); style_run(p.add_run("給教師與高中生的資料解讀、互動儀表板內容與課堂應用"),16,False,INK)
if ASSET.exists():
    pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=pic.add_run(); shape=run.add_picture(str(ASSET),width=Inches(7.0)); shape._inline.docPr.set("descr","學生在教師引導下進行科學實驗與數位學習")
add_para(doc,"OECD PISA 2025 Results (Volume I): Future-Ready Students",11,True,BLUE,before=12,after=3)
add_para(doc,"資料發布：2026 年 9 月 8 日　｜　本文件整理：2026 年 9 月 11 日",10,False,"587080",after=2)
add_para(doc,"臺灣在 OECD 報告中以 Chinese Taipei 呈現。",10,False,"587080")
doc.add_page_break()

add_heading(doc,"先看結論",1)
add_para(doc,"臺灣學生在科學、數學、閱讀三大領域均位居前十，計算問題解決（Computational problem solving）也高於 OECD 平均。高表現的另一面，是成長型思維、學習規劃與自我檢查理解的比例偏低；環境科學能力與行動效能感強，但對環境相關職涯的興趣低於 OECD 平均。",12,False,INK,after=12)
add_table(doc,["觀察面向","臺灣訊號","教學上可追問"],[
    ["學習表現","三大領域均進前十；數學頂尖者 32%","高成就學生是否也能解釋策略、遷移與查證？"],
    ["學習策略","成長型思維 46%；自我提問 33%","是否把規劃、監控與反思明確教出來？"],
    ["數位與 AI","40% 每週用 AI 學習；66% 在課堂學評估 AI 資訊","能否要求學生保留來源、查證並說明修改理由？"],
    ["環境素養","87% 達環境科學基準；職涯興趣 50%","能否把環境議題連到真實工作角色與決策？"],
],[1.25,2.2,3.25])
add_para(doc,"注意：上述「教學上可追問」是本文件根據 OECD 數據提出的教育解讀，並非 OECD 官方政策建議。",10,False,"587080")

add_heading(doc,"1. 受測對象與資料品質",1)
add_table(doc,["項目","臺灣資料"],[["受測學生","7,279 人"],["學校","292 所"],["代表母體","約 18.1 萬名 15 歲學生"],["估計涵蓋率","92%"],["品質判定","符合 PISA 技術品質標準，適合發布"]],[2.1,4.7])
add_para(doc,"學生進行約兩小時測驗，每位學生測兩個領域，並完成約 35 分鐘的背景問卷。校長另填答學校管理、組織與學習環境問卷。",11)

add_heading(doc,"2. 學習表現：高於 OECD 平均",1)
doc.add_picture(str(CHARTS/"scores.png"),width=Inches(7.0)); add_para(doc,"圖 1　平均分數比較。彩色為臺灣，灰色為 OECD 平均。",9,False,"587080",align=WD_ALIGN_PARAGRAPH.CENTER)
add_table(doc,["領域","臺灣","OECD 平均","臺灣概況"],[["科學 Science","540","482","與澳門、日本平均分數無顯著差異"],["數學 Mathematics","546","463","約排名第 4"],["閱讀 Reading","508","461","約排名第 3"],["計算問題解決 CPS","551","500","約排名第 5"]],[2.5,1.0,1.15,2.2])
add_para(doc,"2022 至 2025 年間，臺灣科學約增加 2 分、數學約減少 1 分、閱讀約減少 7 分，均未達統計顯著。不可把小幅變動直接解讀為教育品質上升或下降。",10,False,"587080")

add_heading(doc,"3. 基礎能力與頂尖表現",1)
doc.add_picture(str(CHARTS/"baseline.png"),width=Inches(7.0)); add_para(doc,"圖 2　達熟練度第 2 級（Proficiency Level 2）以上比例。",9,False,"587080",align=WD_ALIGN_PARAGRAPH.CENTER)
add_table(doc,["領域","臺灣達基準","OECD 達基準","臺灣頂尖","OECD 頂尖"],[["科學","87%","74%","20%","7%"],["數學","85%","65%","32%","8%"],["閱讀","82%","69%","13%","6%"]],[2.0,1.2,1.25,1.2,1.2])
add_para(doc,"數學頂尖表現者比例達 32%，為 OECD 平均 8% 的四倍；同時三大領域的基礎能力覆蓋率均高於 OECD。",11)

add_heading(doc,"4. 公平：社經差距仍然存在",1)
add_para(doc,"臺灣社經優勢學生與弱勢學生的科學平均分數相差 91 分，接近 OECD 約 85 分。社經地位解釋臺灣科學表現變異約 11%，OECD 約 12%。弱勢學生中有 11% 進入國內科學表現前四分位，可視為具學業韌性（Academic resilience）；OECD 平均為 12%。",11)
add_para(doc,"2006 至 2025 的長期資料顯示，臺灣科學、數學與閱讀維持高表現，但公平不能只看平均分數。學校仍需持續辨識資源取得、學習支持與高品質任務是否平均分布。",11)

add_heading(doc,"5. 學習投入與自我調節",1)
doc.add_picture(str(CHARTS/"learning.png"),width=Inches(7.0)); add_para(doc,"圖 3　彩色為臺灣，灰色為 OECD 平均。",9,False,"587080",align=WD_ALIGN_PARAGRAPH.CENTER)
add_table(doc,["指標","臺灣","OECD","解讀"],[["遇難題加倍努力","69%","60%","持續投入較高"],["成長型思維","46%","69%","明顯較低"],["規劃讀書方式","30%","37%","需要策略教學"],["自我提問檢查理解","33%","49%","理解監控較弱"],["向教師求助","83%","77%","求助文化較強"],["向同學求助","90%","82%","同儕支持較強"]],[2.3,1.0,1.0,2.6])
add_para(doc,"目標設定指數每增加 1 單位，臺灣科學表現關聯增加 9 分（OECD 3 分）；求助指數每增加 1 單位，關聯增加 13 分（OECD 6 分）。這些是統計關聯，不能直接視為因果效果。",10,False,"587080")

add_heading(doc,"6. 數位與人工智慧（AI）",1)
add_table(doc,["指標","臺灣","OECD 平均"],[["在校數位學習時間／日","1.8 小時","1.7 小時"],["在校數位休閒時間／日","1.1 小時","1.1 小時"],["科學課常受數位裝置分心","13%","28%"],["每週用 AI 協助學習","40%","46%"],["以 AI 做初步研究","22%","31%"],["以 AI 摘要文本","20%","30%"],["以 AI 撰寫草稿","21%","29%"],["課堂學習評估 AI 生成資訊","65.9%","62.6%"]],[3.5,1.6,1.7])
add_para(doc,"臺灣學生的數位使用時間與 OECD 接近，課堂分心回報較少。OECD 指出，在臺灣，科學課數位分心與科學分數差異並無顯著關聯。這不表示分心沒有影響，而是目前資料不足以支持該關聯。",11)
add_heading(doc,"可直接採用的 AI 任務設計",2)
for text in ["保留提示詞、AI 回答與修改紀錄。","用至少兩個可追溯來源查證三項主張。","標示保留、改寫與刪除內容，並說明判斷依據。","評量重點放在證據品質與推理，而非只看輸出流暢度。"]:
    p=doc.add_paragraph(style="List Bullet"); style_run(p.add_run(text),11)

add_heading(doc,"7. 校園支持、資源與安全",1)
add_table(doc,["指標","臺灣","OECD 平均／趨勢"],[["教師關心每位學生學習","79%","69%"],["教師教到學生理解","72%","66%"],["需要時提供額外協助","84%","73%"],["每月至少數次遭受霸凌","11%","20%"],["測驗前兩週曾遲到","38%","50%"],["教學受師資短缺影響","28%","39%"],["基礎設施短缺影響教學","8%","2022 年臺灣為 19%"]],[3.3,1.5,2.0])
add_para(doc,"家庭支持回報下降：每週詢問孩子在校經驗由 2022 年 62% 降至 2025 年 59%；每週討論在校問題由 58% 降至 53%。OECD 指出這是多數參與教育體系的共同現象。",11)

add_heading(doc,"8. 環境科學與未來職涯",1)
doc.add_picture(str(CHARTS/"environment.png"),width=Inches(7.0)); add_para(doc,"圖 4　環境素養與行動相關指標。",9,False,"587080",align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc,"臺灣學生在環境科學與行動效能感上表現突出：87% 達環境科學基準、92% 相信自己能帶來正向影響、86% 認為自己知道如何與他人合作。然而，只有 50% 對未來從事環境保護相關工作有興趣，低於 OECD 的 62%。",11)
add_heading(doc,"教學連結",2)
add_para(doc,"把環境議題從「知道」推進到「角色與決策」。可用校園能源、飲食或交通為情境，讓學生分別扮演資料分析師、工程師、設計師、政策溝通者，提出可驗證的改善方案。",11)

add_heading(doc,"9. 四個教學行動",1)
add_table(doc,["行動","作法","學生可回答的問題"],[
    ["教會學習策略","任務前規劃、進行中監控、結束後反思","我何時發現自己沒有理解？下一次要改什麼？"],
    ["閱讀加入查證","比較兩則立場不同的文本，區分事實、推論與意見","哪一項證據最能改變我的判斷？"],
    ["AI 從產出轉向評估","保留提示詞，查證並說明修改理由","我保留、修正與拒絕了什麼？為什麼？"],
    ["環境連結職涯","分派跨專業角色，完成真實情境提案","不同職業如何共同解決同一個環境問題？"],
],[1.5,2.7,2.7])

add_heading(doc,"10. 閱讀資料時的限制",1)
for text in ["問卷指標可能受文化規範、作答習慣與題意理解影響。","相關關係不能直接解讀為因果；控制社經背景也不等於完成因果識別。","平均分數與排名具有抽樣不確定性；分數接近時未必存在顯著差異。","PISA 評估的是 15 歲學生面對真實情境的知識應用，不等同學校課程考試。","英語外語評量與數位自我導向學習的完整分析預計 2027 年發布。"]:
    p=doc.add_paragraph(style="List Bullet"); style_run(p.add_run(text),11)

add_heading(doc,"資料來源",1)
sources=[
    "OECD (2026), PISA 2025 Results (Volume I): Future-Ready Students. DOI: 10.1787/73451bc5-en",
    "OECD (2026), PISA 2025 Results (Volume I): Chinese Taipei - Country Note.",
    "OECD Education GPS: Chinese Taipei - Student performance (PISA 2025).",
]
for s in sources: add_para(doc,s,10,False,BLUE,after=5)
add_para(doc,"授權與改編說明：OECD 原始內容依 Creative Commons Attribution 4.0 International（CC BY 4.0）提供。本文件為繁體中文改編與教學重整；若與原文有出入，以英文原文為準。本文件中的教學解讀不代表 OECD 或其會員國的官方立場。",9,False,"587080",before=10)

# Page numbers in footer
for section in doc.sections:
    f=section.footer; p=f.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("PISA 2025 臺灣學習羅盤  |  "),9,False,"587080")
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); p._p.append(fld)

doc.core_properties.title="PISA 2025 臺灣學習羅盤"
doc.core_properties.subject="PISA 2025 臺灣資料解讀與教師教學應用"
doc.core_properties.author=""
doc.save(OUT)
print(OUT)
